"""Integration tests for all required document formats and error isolation."""

from pathlib import Path

import docx
import pptx
from PIL import Image

from core.converter import ConversionManager
from core.image_processor import ImageProcessor


def _create_minimal_pdf(file_path: Path, text: str = "Sample Digital PDF Document") -> None:
    stream = f"BT /F1 12 Tf 50 750 Td ({text}) Tj ET"
    content = (
        "%PDF-1.4\n"
        "1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
        "2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
        "3 0 obj<</Type/Page/MediaBox[0 0 595 842]/Parent 2 0 R"
        "/Resources<</Font<</F1 4 0 R>>>>/Contents 5 0 R>>endobj\n"
        "4 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj\n"
        f"5 0 obj<</Length {len(stream)}>>stream\n{stream}\nendstream\nendobj\n"
        "xref\n0 6\n"
        "0000000000 65535 f \n"
        "0000000009 00000 n \n"
        "0000000056 00000 n \n"
        "0000000111 00000 n \n"
        "0000000234 00000 n \n"
        "0000000305 00000 n \n"
        "trailer<</Size 6/Root 1 0 R>>\nstartxref\n380\n%%EOF\n"
    )
    file_path.write_bytes(content.encode("latin-1"))


def test_pdf_digital_conversion(tmp_path):
    pdf_file = tmp_path / "digital_sample.pdf"
    _create_minimal_pdf(pdf_file, "Antigravity PDF Digital Extraction")

    mgr = ConversionManager()
    res = mgr.convert_file(pdf_file, engine_choice="markitdown", output_dir=tmp_path / "out")

    assert res.success is True
    assert "Antigravity PDF Digital Extraction" in res.markdown_content
    assert res.output_md_path.is_file()


def test_docx_conversion_with_headings_and_tables(tmp_path):
    doc = docx.Document()
    doc.add_heading("Docx Heading 1", level=1)
    doc.add_paragraph("First paragraph content.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Col A"
    table.cell(0, 1).text = "Col B"
    table.cell(1, 0).text = "Data A"
    table.cell(1, 1).text = "Data B"

    docx_file = tmp_path / "sample_doc.docx"
    doc.save(docx_file)

    mgr = ConversionManager()
    res = mgr.convert_file(docx_file, engine_choice="markitdown", output_dir=tmp_path / "out")

    assert res.success is True
    assert "Docx Heading 1" in res.markdown_content
    assert "First paragraph content" in res.markdown_content
    assert "Col A" in res.markdown_content
    assert "Data B" in res.markdown_content


def test_pptx_conversion(tmp_path):
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Universal Slide Presentation"

    pptx_file = tmp_path / "sample_pres.pptx"
    prs.save(pptx_file)

    mgr = ConversionManager()
    res = mgr.convert_file(pptx_file, engine_choice="markitdown", output_dir=tmp_path / "out")

    assert res.success is True
    assert "Universal Slide Presentation" in res.markdown_content


def test_batch_conversion_error_isolation(tmp_path):
    # Create 1 valid CSV
    csv1 = tmp_path / "good.csv"
    csv1.write_text("id,name\n1,Alice\n", encoding="utf-8")

    # Create 1 corrupt PDF
    bad_pdf = tmp_path / "bad.pdf"
    bad_pdf.write_bytes(b"NOT_A_VALID_PDF")

    # Create 1 valid PDF
    good_pdf = tmp_path / "good.pdf"
    _create_minimal_pdf(good_pdf, "Valid Batch Document")

    mgr = ConversionManager()
    results = mgr.convert_batch(
        [csv1, bad_pdf, good_pdf],
        engine_choice="markitdown",
        output_dir=tmp_path / "out",
    )

    assert len(results) == 3
    # 1st: success
    assert results[0].success is True
    assert results[0].filename == "good.csv"
    # 2nd: isolated failure
    assert results[1].success is False
    assert results[1].filename == "bad.pdf"
    # 3rd: success, unaffected by the failure of 2nd file!
    assert results[2].success is True
    assert results[2].filename == "good.pdf"


def test_image_processor_and_reference(tmp_path):
    out_dir = tmp_path / "out"
    processor = ImageProcessor(base_output_dir=out_dir, doc_stem="manual")

    img = Image.new("RGB", (60, 60), color="blue")
    rel_path = processor.save_image(img, image_format="PNG", caption="Blue Square")

    assert rel_path is not None
    assert "manual_assets/image_001.png" in rel_path

    tag = processor.get_markdown_tag(rel_path, caption="Blue Square")
    assert tag == f"![Blue Square]({rel_path})"

    # Verify physical file existence
    saved_file = out_dir / rel_path
    assert saved_file.is_file()
    assert saved_file.stat().st_size > 0
