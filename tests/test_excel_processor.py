"""Unit tests for Excel processor."""

from datetime import date

import openpyxl

from core.excel_processor import ExcelProcessor


def test_excel_processor_multi_sheet(tmp_path):
    wb = openpyxl.Workbook()

    # Sheet 1: Sales
    ws1 = wb.active
    ws1.title = "Sales"
    ws1.append(["Product", "Qty", "Price", "Date"])
    ws1.append(["Laptop", 5, 1200.50, date(2026, 1, 15)])
    ws1.append(["Mouse", 10, 25.00, date(2026, 1, 16)])

    # Sheet 2: Merged Cells
    ws2 = wb.create_sheet(title="Departments")
    ws2.append(["Region", "Department", "Staff"])
    ws2.append(["West", "Engineering", 20])
    ws2.append(["West", "Marketing", 15])
    # Merge region column cells A2:A3
    ws2.merge_cells("A2:A3")

    # Sheet 3: Empty Sheet
    wb.create_sheet(title="EmptySheet")

    excel_file = tmp_path / "test_report.xlsx"
    wb.save(excel_file)
    wb.close()

    markdown_res, metadata = ExcelProcessor.process_file(excel_file)

    # Assertions
    assert "## Worksheet: Sales" in markdown_res
    assert "## Worksheet: Departments" in markdown_res
    assert "## Worksheet: EmptySheet" in markdown_res
    assert "Laptop" in markdown_res
    assert "Mouse" in markdown_res
    assert "Engineering" in markdown_res
    assert "Marketing" in markdown_res

    assert metadata["total_sheets"] == 3
    assert len(metadata["sheets"]) == 3
    assert metadata["sheets"][0]["sheet_name"] == "Sales"
    assert metadata["sheets"][0]["rows"] == 3
    assert metadata["sheets"][0]["columns"] == 4
    assert metadata["sheets"][2]["is_empty"] is True

